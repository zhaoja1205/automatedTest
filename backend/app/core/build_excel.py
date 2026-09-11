#!/usr/bin/env python3
"""Fill the SWE1 requirements-table Excel template from a project JSON file.

Usage:
    python3 build_excel.py --project project.json --output "ProjectX_Detailed Software Requirements Table需求详细表.xlsx"

The project JSON schema (see references/data_model.md):
{
  "project": {...}, "document": {...},
  "inputs": [...], "requirements": [...], "risks": [...]
}
"""
import argparse
import copy
import datetime
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

TEMPLATE = Path(__file__).resolve().parent / "aspice_templates" / "requirements_table_template.xlsx"

# Requirement sheet: JSON field -> column letter
REQ_COLUMNS = {
    "input": "A", "chapter": "B", "no": "C", "content": "D", "flag": "E",
    "or_id": "F", "or_desc": "G", "req_id": "H", "sw_req_desc": "I",
    "test_case_id": "J", "module": "K", "category": "L", "asil": "M",
    "correctness": "N", "feasibility": "O", "exception": "P", "priority": "Q",
    "release_time": "R", "ra_deadline": "S", "actual_time": "T",
    "release_version": "U", "owner": "V", "memo": "W", "arch_doc": "X",
}

DATE_FIELDS = ("release_time", "ra_deadline", "actual_time")
DATE_COLS = {REQ_COLUMNS[f] for f in DATE_FIELDS}


def parse_date(val):
    """Parse a date value; returns datetime or None. N/A-style strings stay None."""
    if val in (None, "", "N/A", "NA", "TBD", "/"):
        return None
    if isinstance(val, (datetime.date, datetime.datetime)):
        return val
    s = str(val).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S"):
        try:
            return datetime.datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def as_text(val):
    """List values joined by newline (Excel multi-line cell), others str()-able."""
    if isinstance(val, list):
        return "\n".join(str(v) for v in val)
    return val


def precheck(proj):
    """Fail fast with a readable message instead of a KeyError traceback."""
    pr, pj = proj.get("project"), proj.get("document")
    problems = []
    if not isinstance(pr, dict) or not pr.get("name"):
        problems.append("project.name is required")
    if not isinstance(pj, dict) or not pj.get("doc_number"):
        problems.append("document.doc_number is required")
    if problems:
        raise SystemExit("ERROR: invalid project JSON:\n  - " + "\n  - ".join(problems))


def fill_cover(ws, proj):
    doc_num = proj.get("document", {}).get("doc_number", "")
    ws["A3"] = f"文件编号：{doc_num}" if doc_num else "文件编号："
    ws["A15"] = f"{proj['project']['name']} Detailed Software Requirements Table需求详细表"
    ws["A17"] = f"Version/版本：{proj['document'].get('version', 'V1.0')}"


def fill_doc_record(ws, proj):
    """Document record sheet: one row per version entry, data rows 4..14."""
    records = proj.get("document", {}).get("history", [])
    for i, rec in enumerate(records):
        r = 4 + i
        if r > 14:
            print(f"[WARN] doc record overflow: {len(records)} records > 11 rows, extras dropped", file=sys.stderr)
            break
        ws.cell(row=r, column=2, value=rec.get("version", ""))
        ws.cell(row=r, column=3, value=rec.get("description", ""))
        dcell = ws.cell(row=r, column=4, value=parse_date(rec.get("date")))
        if isinstance(dcell.value, (datetime.date, datetime.datetime)):
            dcell.number_format = "yyyy/m/d"
        ws.cell(row=r, column=5, value=rec.get("author", ""))
        ws.cell(row=r, column=6, value=rec.get("reviewers", ""))
        ws.cell(row=r, column=7, value=rec.get("approver", "/"))


def fill_input(ws, proj):
    ws["C2"] = proj["project"].get("module_name", "")
    inputs = proj.get("inputs", [])
    for i, inp in enumerate(inputs):
        r = 6 + i
        if r > 15:
            print(f"[WARN] input list overflow: {len(inputs)} docs > 10 rows, extras dropped", file=sys.stderr)
            break
        name = inp.get("name", "") if isinstance(inp, dict) else str(inp)
        ws.cell(row=r, column=3, value=name)


def fill_requirement(ws, proj):
    """Requirement sheet: data starts at row 3; row 3 carries the base cell style."""
    reqs = proj.get("requirements", [])
    if not reqs:
        print("[WARN] no requirements in project JSON; requirement sheet left empty", file=sys.stderr)
        return

    # capture base style from template row 3 (cleared but styled)
    base = 3
    style_src = {}
    for c in range(1, 25):
        cell = ws.cell(row=base, column=c)
        style_src[c] = (copy.copy(cell.font), copy.copy(cell.border),
                        copy.copy(cell.fill), copy.copy(cell.alignment),
                        cell.number_format)

    for i, req in enumerate(reqs):
        r = base + i
        for field, col in REQ_COLUMNS.items():
            val = req.get(field)
            if field in DATE_FIELDS:
                d = parse_date(val)
                cell = ws.cell(row=r, column=column_index_from_string(col))
                if d is not None:
                    cell.value = d
                else:
                    cell.value = as_text(val)
            elif field == "priority":
                try:
                    val = int(val)
                except (TypeError, ValueError):
                    pass
                ws[f"{col}{r}"] = val
            else:
                ws[f"{col}{r}"] = as_text(val)
        # apply row style
        for c in range(1, 25):
            cell = ws.cell(row=r, column=c)
            f, b, fl, al, nf = style_src[c]
            cell.font, cell.border, cell.fill, cell.alignment = f, b, fl, al
        # date format on date columns
        for col in DATE_COLS:
            cell = ws[f"{col}{r}"]
            if isinstance(cell.value, (datetime.date, datetime.datetime)):
                cell.number_format = "yyyy/m/d"
            else:
                cell.number_format = "General"  # N/A / NA text must not carry a date format


def fill_risk(ws, proj):
    risks = proj.get("risks", [])
    for i, risk in enumerate(risks):
        r = 4 + i
        if r > 7:
            print(f"[WARN] risk overflow: {len(risks)} risks > 4 rows, extras dropped", file=sys.stderr)
            break
        ws.cell(row=r, column=2, value=risk.get("req_id", ""))
        ws.cell(row=r, column=3, value=risk.get("function", ""))
        ws.cell(row=r, column=4, value=risk.get("risk", ""))
        ws.cell(row=r, column=5, value=risk.get("solution", ""))
        ws.cell(row=r, column=6, value=risk.get("owner", ""))


def build(project: dict, output: str, template: str = None) -> str:
    """Build Excel from a project dict (in-memory), for programmatic use.

    Args:
        project: project dict following the skill data_model schema
        output: output xlsx path
        template: optional template path (defaults to bundled)
    Returns:
        output path
    """
    precheck(project)
    tpl = template or str(TEMPLATE)
    wb = openpyxl.load_workbook(tpl)
    fill_cover(wb["Cover 文件封面"], project)
    fill_doc_record(wb["Document record文件修改控制"], project)
    fill_input(wb["Input 输入"], project)
    fill_requirement(wb["Requirement 需求"], project)
    fill_risk(wb["Risk 风险"], project)
    wb.save(output)
    return output


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--project", required=True, help="project JSON file")
    ap.add_argument("--output", required=True, help="output xlsx path")
    ap.add_argument("--template", default=str(TEMPLATE), help="template xlsx (default: bundled)")
    args = ap.parse_args()

    with open(args.project, encoding="utf-8") as f:
        proj = json.load(f)
    build(proj, args.output, args.template)
    print(f"OK: {args.output}")
    print(f"    requirements: {len(proj.get('requirements', []))}, "
          f"inputs: {len(proj.get('inputs', []))}, risks: {len(proj.get('risks', []))}")


if __name__ == "__main__":
    main()
