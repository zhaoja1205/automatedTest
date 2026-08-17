"""
Excel 用例表加载与结果回写。

与原 PyQt 版保持兼容的列映射与动态表头识别。
"""
import os
from dataclasses import asdict
from datetime import datetime
from typing import List, Optional, Dict, Any
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from app.core.test_case import TestCase


class ExcelHandler:
    """
    Excel 工作簿的加载、布局探测、用例构建与结果保存。
    """

    @staticmethod
    def is_supported_test_sheet(sheet_name: str) -> bool:
        """判断工作表是否视为「用例表」"""
        if not sheet_name or not str(sheet_name).strip():
            return False
        n = str(sheet_name).strip()
        nl = n.lower()
        if "功能" in n or "functional" in nl:
            return True
        if "故障" in n or "fault" in nl:
            return True
        if "稳定" in n or "stability" in nl:
            return True
        return False

    @staticmethod
    def sheet_category(sheet_name: str) -> Optional[str]:
        """将工作表名归类到界面套件：functional / fault / stability"""
        if not sheet_name or not str(sheet_name).strip():
            return None
        s = str(sheet_name).strip()
        sl = s.lower()
        if "功能" in s or "functional" in sl:
            return "functional"
        if "故障" in s or "fault" in sl:
            return "fault"
        if "稳定" in s or "stability" in sl:
            return "stability"
        return None

    @staticmethod
    def ordered_test_sheet_names(workbook: Workbook) -> List[str]:
        """按 功能 → 故障 → 稳定 → 其它已识别 的顺序排列工作表名"""
        names = [x for x in workbook.sheetnames if ExcelHandler.is_supported_test_sheet(x)]

        def sort_key(name: str):
            s = name.strip()
            sl = s.lower()
            if "功能" in s or "functional" in sl:
                return (0, s)
            if "故障" in s or "fault" in sl:
                return (1, s)
            if "稳定" in s or "stability" in sl:
                return (2, s)
            return (3, s)

        return sorted(names, key=sort_key)

    # 默认列号兜底（1-based）
    COLUMN_MAP = {
        'case_id': 2,
        'test_type': 3,
        'requirement_id': 4,
        'design_method': 5,
        'description': 6,
        'prerequisites': 7,
        'test_steps': 8,
        'expected_result': 9,
        'priority': 10,
    }

    RESULT_COLUMN_MAP = {
        'actual_result': 12,
        'bug_id': 13,
        'remarks': 14,
        'test_version': 15,
        'tester': 16,
        'test_date': 17,
    }

    # 表头关键词
    HEADER_KEYWORDS = {
        'case_id': ['测试用例id', 'test case id', '用例id', 'case id', '用例编号'],
        'test_type': ['测试类型', 'test type'],
        'requirement_id': ['需求id', 'requirement id', '需求编号'],
        'design_method': ['设计方法', 'design method', '用例设计'],
        'description': ['测试用例描述', 'test case desc', '用例描述', 'description'],
        'prerequisites': ['前置条件', 'prerequisites', 'precondition', 'pre-condition'],
        'test_steps': ['测试步骤', 'testing steps', 'test steps', '步骤'],
        'expected_result': ['预期结果', 'expected result', 'expected results'],
        'priority': ['优先级', 'priority'],
    }

    RESULT_HEADER_KEYWORDS = {
        'actual_result': ['实际结果', 'actual result', 'actual results'],
        'bug_id': ['bug', 'defect', 'bug id'],
        'remarks': ['备注', 'remark', 'note', '测试结果的log'],
        'test_version': ['测试版本', 'test version'],
        'tester': ['实施者', 'implementer', 'tester', '测试人员'],
        'test_date': ['测试日期', 'date', 'test date'],
    }

    HEADER_ROW = 13
    DATA_START_ROW = 15

    STATUS_COLORS = {
        'Pass': '90EE90',
        'Fail': 'FF6B6B',
        'NT': 'FFFFE0',
        'BLOCK': 'FFA500',
        'NA': 'D3D3D3',
    }

    def __init__(self, file_path: str = None):
        self.file_path = file_path
        self.workbook: Optional[Workbook] = None
        self.test_cases: List[TestCase] = []

    @classmethod
    def _cell_text(cls, ws, row: int, col: int) -> str:
        """读取单元格并规范化为用于匹配的文本"""
        raw = ws.cell(row=row, column=col).value
        if not raw:
            return ""
        return str(raw).replace('\n', ' ').replace('\r', '').strip().lower()

    @classmethod
    def _count_header_matches(cls, ws, row: int, all_keywords: Dict[str, List[str]]) -> int:
        """计算某行中命中关键词的字段个数"""
        max_col = ws.max_column or 30
        matched_fields = set()
        for col in range(1, max_col + 1):
            cell = cls._cell_text(ws, row, col)
            if not cell:
                continue
            for field, kws in all_keywords.items():
                if field in matched_fields:
                    continue
                for kw in kws:
                    if kw.lower() in cell:
                        matched_fields.add(field)
                        break
        return len(matched_fields)

    @classmethod
    def _find_header_and_data_rows(
        cls, ws, all_keywords: Dict[str, List[str]], scan_rows: int = 25, min_matches: int = 2
    ) -> tuple:
        """定位表头行和数据起始行"""
        best_row = None
        best_count = min_matches - 1

        for row in range(1, scan_rows + 1):
            cnt = cls._count_header_matches(ws, row, all_keywords)
            if cnt > best_count:
                best_count = cnt
                best_row = row

        if best_row is None:
            return None, None

        data_start = best_row + 1
        for r in range(best_row + 1, best_row + 5):
            if r > ws.max_row:
                break
            max_col = ws.max_column or 30
            non_empty = sum(
                1 for c in range(1, max_col + 1)
                if ws.cell(row=r, column=c).value not in (None, "")
            )
            if non_empty >= 2:
                data_start = r
                break

        return best_row, data_start

    @classmethod
    def _detect_column_map(
        cls, ws, header_keywords: Dict[str, List[str]], default_map: Dict[str, int],
        header_row: Optional[int] = None, extra_scan_rows: int = 5,
    ) -> Dict[str, int]:
        """在表头区域扫描单元格，将列标题映射为字段名"""
        detected: Dict[str, int] = {}
        start = header_row if header_row else 1
        end = start + extra_scan_rows

        for row in range(start, end + 1):
            if row > (ws.max_row or 9999):
                break
            max_col = ws.max_column or 30
            for col in range(1, max_col + 1):
                cell = cls._cell_text(ws, row, col)
                if not cell:
                    continue
                for field, keywords in header_keywords.items():
                    if field in detected:
                        continue
                    for kw in keywords:
                        if kw.lower() == cell or kw.lower() in cell:
                            detected[field] = col
                            break

        result = dict(default_map)
        result.update(detected)
        return result

    CONTENT_FEATURES = {
        'test_steps': [
            'sudo ', '执行指令', '执行命令', '输入命令', './nvsipl', 'nvsipl_camera',
            'i2cmastercmd', 'ssh登录', 'ssh nvidia', '1、执行', '1.执行',
        ],
        'prerequisites': [
            '硬件正确连接', '硬件设备正确', '将驱动文件', '将测试so', 'ssh登录',
            '前提：', '前提条件', '先执行',
        ],
        'expected_result': [
            'fps稳定输出', '无报错', '正常输出', '无异常', '稳定在', 'pass',
            '可生成', '正确抓出',
        ],
        'description': [
            '出图', '起流', '拍照', '故障诊断', '解串器', '加串器', '传感器', '帧率',
        ],
    }

    @classmethod
    def _infer_columns_by_content(cls, ws, data_start_row: int) -> Dict[str, int]:
        """无表头时依据数据格文本推断列语义"""
        max_col = ws.max_column or 15
        field_scores: Dict[str, Dict[int, int]] = {
            field: {} for field in cls.CONTENT_FEATURES
        }
        sample_rows = min(data_start_row + 10, (ws.max_row or 9999) + 1)
        for row in range(data_start_row, sample_rows):
            for col in range(1, max_col + 1):
                cell = cls._cell_text(ws, row, col)
                if not cell:
                    continue
                for field, kws in cls.CONTENT_FEATURES.items():
                    for kw in kws:
                        if kw.lower() in cell:
                            field_scores[field][col] = field_scores[field].get(col, 0) + 1
                            break

        inferred: Dict[str, int] = {}
        for field, scores in field_scores.items():
            if not scores:
                continue
            best_col = max(scores, key=lambda c: scores[c])
            if scores[best_col] >= 2:
                inferred[field] = best_col
        return inferred

    def _detect_sheet_layout(self, ws) -> tuple:
        """单张 Sheet 的完整布局解析"""
        all_kw = {**self.HEADER_KEYWORDS, **self.RESULT_HEADER_KEYWORDS}
        header_row, data_start_row = self._find_header_and_data_rows(ws, all_kw)

        if header_row is not None:
            col_map = self._detect_column_map(
                ws, self.HEADER_KEYWORDS, self.COLUMN_MAP,
                header_row=header_row, extra_scan_rows=2,
            )
            result_col_map = self._detect_column_map(
                ws, self.RESULT_HEADER_KEYWORDS, self.RESULT_COLUMN_MAP,
                header_row=header_row, extra_scan_rows=3,
            )
            return header_row, data_start_row, col_map, result_col_map

        inferred_data_start = self.DATA_START_ROW
        for r in range(1, (ws.max_row or 30) + 1):
            max_col = ws.max_column or 15
            non_empty = sum(
                1 for c in range(1, max_col + 1)
                if ws.cell(row=r, column=c).value not in (None, "")
            )
            if non_empty >= 3:
                inferred_data_start = r
                break

        inferred = self._infer_columns_by_content(ws, inferred_data_start)
        col_map = dict(self.COLUMN_MAP)
        col_map.update(inferred)

        return None, inferred_data_start, col_map, dict(self.RESULT_COLUMN_MAP)

    def load(self, file_path: str = None) -> List[TestCase]:
        """打开 Excel，按「功能 / 故障 / 稳定」类 Sheet 顺序解析用例"""
        if file_path:
            self.file_path = file_path

        if not self.file_path or not os.path.exists(self.file_path):
            raise FileNotFoundError(f"文件不存在: {self.file_path}")

        if self.workbook:
            try:
                self.workbook.close()
            except Exception:
                pass
            self.workbook = None

        self.workbook = load_workbook(self.file_path, data_only=False, read_only=False)
        self.test_cases = []

        sheet_names = self.ordered_test_sheet_names(self.workbook)
        if not sheet_names:
            raise ValueError(
                "未找到用例工作表。请确认存在名称中含「功能」「故障」「稳定」"
                "或 Functional / Fault / Stability 的 Sheet。"
            )

        if not hasattr(self, '_sheet_col_maps'):
            self._sheet_col_maps: Dict[str, tuple] = {}

        for sheet_name in sheet_names:
            ws = self.workbook[sheet_name]

            header_row, data_start_row, col_map, result_col_map = self._detect_sheet_layout(ws)
            self._sheet_col_maps[sheet_name] = (col_map, result_col_map)

            cat = self.sheet_category(sheet_name)
            default_selected = cat == "functional"
            seq = 0

            for row_idx in range(data_start_row, (ws.max_row or 0) + 1):
                case_id = self._get_cell_value(ws, row_idx, col_map['case_id'])

                max_col = ws.max_column or 30
                row_has_content = any(
                    ws.cell(row=row_idx, column=c).value not in (None, "")
                    for c in range(1, max_col + 1)
                )
                if not row_has_content:
                    continue

                if not case_id:
                    steps = self._get_cell_value(ws, row_idx, col_map['test_steps'])
                    expected = self._get_cell_value(ws, row_idx, col_map['expected_result'])
                    prerequisites = self._get_cell_value(ws, row_idx, col_map['prerequisites'])
                    if not steps and not expected and not prerequisites:
                        continue
                    seq += 1
                    case_id = f"{sheet_name[:4]}_{row_idx:03d}"
                else:
                    seq += 1

                tc = TestCase(
                    case_id=case_id,
                    case_key=f"{sheet_name}:{row_idx}:{case_id}",
                    test_type=self._get_cell_value(ws, row_idx, col_map['test_type']),
                    requirement_id=self._get_cell_value(ws, row_idx, col_map['requirement_id']),
                    design_method=self._get_cell_value(ws, row_idx, col_map['design_method']),
                    description=self._get_cell_value(ws, row_idx, col_map['description']),
                    prerequisites=self._get_cell_value(ws, row_idx, col_map['prerequisites']),
                    test_steps=self._get_cell_value(ws, row_idx, col_map['test_steps']),
                    expected_result=self._get_cell_value(ws, row_idx, col_map['expected_result']),
                    priority=self._get_cell_value(ws, row_idx, col_map['priority']),
                    row_number=row_idx,
                    source_sheet=sheet_name,
                    selected=default_selected,
                    status="NT"
                )
                tc.expected_keywords = self._parse_expected_keywords(tc.expected_result)
                self.test_cases.append(tc)

        return self.test_cases

    def _get_cell_value(self, ws, row: int, col: int) -> str:
        """读取单元格原始值并转为字符串"""
        cell = ws.cell(row=row, column=col)
        return str(cell.value).strip() if cell.value else ""

    @staticmethod
    def _parse_expected_keywords(expected_text: str) -> List[str]:
        """解析预期结果文本，填充 expected_keywords 供前端预览与执行器匹配使用"""
        if not expected_text:
            return []
        try:
            from app.core.expected_parser import ExpectedResultParser
            criteria = ExpectedResultParser().parse(expected_text)
            return criteria.keywords
        except Exception:
            return []

    def save_results(self, test_cases: List[TestCase] = None, output_path: str = None):
        """将内存中用例的执行结果写回工作簿并保存"""
        if test_cases is None:
            test_cases = self.test_cases

        if not self.workbook:
            raise ValueError("未加载 Excel 文件")

        def _worksheet_for_case(tc: TestCase):
            name = (tc.source_sheet or "").strip()
            if name and name in self.workbook.sheetnames:
                return self.workbook[name]
            for cand in self.workbook.sheetnames:
                if "功能" in cand or "Functional" in cand.lower():
                    return self.workbook[cand]
            return self.workbook[self.workbook.sheetnames[0]]

        def _result_col_for(tc: TestCase) -> Dict[str, int]:
            sheet_maps = getattr(self, '_sheet_col_maps', {})
            entry = sheet_maps.get(tc.source_sheet or "")
            if entry:
                return entry[1]
            return self.RESULT_COLUMN_MAP

        for tc in test_cases:
            if tc.row_number <= 0:
                continue

            row = tc.row_number
            ws = _worksheet_for_case(tc)
            rcm = _result_col_for(tc)

            cell = ws.cell(row=row, column=rcm['actual_result'])
            cell.value = tc.status

            if tc.status in self.STATUS_COLORS:
                cell.fill = PatternFill(
                    start_color=self.STATUS_COLORS[tc.status],
                    end_color=self.STATUS_COLORS[tc.status],
                    fill_type='solid'
                )

            remark_value = tc.remarks or ""
            if not remark_value and tc.log_file:
                remark_value = tc.log_file
            if remark_value:
                ws.cell(row=row, column=rcm['remarks']).value = remark_value

            if tc.test_version:
                ws.cell(row=row, column=rcm['test_version']).value = tc.test_version

            if tc.tester:
                ws.cell(row=row, column=rcm['tester']).value = tc.tester

            if tc.test_date:
                ws.cell(row=row, column=rcm['test_date']).value = tc.test_date

        save_path = output_path or self.file_path
        self.workbook.save(save_path)
        return save_path

    def get_statistics(self, test_cases: List[TestCase] = None) -> Dict[str, int]:
        """统计用例数量"""
        if test_cases is None:
            test_cases = self.test_cases

        return {
            'total': len(test_cases),
            'selected': sum(1 for tc in test_cases if tc.selected),
            'pass': sum(1 for tc in test_cases if tc.status == 'Pass'),
            'fail': sum(1 for tc in test_cases if tc.status == 'Fail'),
            'nt': sum(1 for tc in test_cases if tc.status == 'NT'),
            'block': sum(1 for tc in test_cases if tc.status == 'BLOCK'),
            'na': sum(1 for tc in test_cases if tc.status == 'NA'),
        }

    def close(self):
        if self.workbook:
            self.workbook.close()
            self.workbook = None