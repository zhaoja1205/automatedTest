"""
ASPICE 文档模块 API 路由。

提供 SWE.1 需求分析 / SWE.2 架构设计的生成、保存、导出接口。
与 creator 模块共享同一项目实体（runtime/creator_projects/<id>/project.json）。
"""
from fastapi import APIRouter, HTTPException, Request, Query, UploadFile, File
from fastapi.responses import FileResponse
from pydantic import BaseModel
from typing import Optional, List
import os

from app.core import creator_store
from app.core import id_chain
from app.core import req_generator
from app.core import arch_generator
from app.core.config_store import ConfigStore
from app.core.test_case import AIConfig

router = APIRouter()

_global_config_store = ConfigStore(base_dir="runtime")


def _get_ai_service():
    from app.ai.service import AIService
    ai_config = _global_config_store.load("ai_config", AIConfig, AIConfig())
    return AIService(ai_config.model_dump())


# ---- 请求模型 ----

class CreateAspiceProjectRequest(BaseModel):
    name: str


class SaveSwe1Request(BaseModel):
    project_code: Optional[str] = None
    requirements: Optional[list] = None
    topology: Optional[list] = None
    kpi: Optional[list] = None
    current_step: Optional[int] = None


class SaveSwe2Request(BaseModel):
    mappings: Optional[list] = None
    components: Optional[list] = None
    current_step: Optional[int] = None


class SaveAspiceRequest(BaseModel):
    project_code: Optional[str] = None
    swe1: Optional[dict] = None
    swe2: Optional[dict] = None


class GenerateReqRequest(BaseModel):
    raw_text: str
    use_ai: bool = True


# ---- 项目级 ----

@router.get("/projects")
async def list_aspice_projects():
    """列出 ASPICE 文档项目。"""
    return creator_store.list_projects(project_type="aspice")


@router.post("/projects")
async def create_aspice_project(req: CreateAspiceProjectRequest):
    """创建 ASPICE 文档项目。"""
    if not req.name.strip():
        raise HTTPException(status_code=400, detail="项目名称不能为空")
    return creator_store.create_project(req.name.strip(), project_type="aspice")


@router.get("/projects/{project_id}")
async def get_aspice_project(project_id: str):
    """获取项目的 ASPICE 数据段。"""
    project = creator_store.get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    aspice = project.get("aspice")
    if not aspice:
        aspice = creator_store._default_aspice()
        project["aspice"] = aspice
        creator_store.save_project(project)
    return aspice


@router.put("/projects/{project_id}")
async def save_aspice_project(project_id: str, req: SaveAspiceRequest):
    """一次性保存完整 ASPICE 数据段，避免 SWE.1/SWE.2 分两次落盘。"""
    project = creator_store.get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    aspice = project.get("aspice") or creator_store._default_aspice()
    update = req.model_dump(exclude_none=True)
    if "project_code" in update:
        aspice["project_code"] = update["project_code"]
    if "swe1" in update:
        aspice["swe1"] = {**aspice.get("swe1", {}), **update["swe1"]}
    if "swe2" in update:
        aspice["swe2"] = {**aspice.get("swe2", {}), **update["swe2"]}
    project["aspice"] = aspice
    creator_store.save_project(project)
    return {"message": "已保存"}


@router.put("/projects/{project_id}/swe1")
async def save_swe1(project_id: str, req: SaveSwe1Request):
    """保存 SWE.1 数据。"""
    project = creator_store.get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    aspice = project.get("aspice") or creator_store._default_aspice()
    swe1 = aspice.get("swe1", {})
    update = req.model_dump(exclude_none=True)
    if "project_code" in update:
        aspice["project_code"] = update.pop("project_code")
    for k, v in update.items():
        swe1[k] = v
    aspice["swe1"] = swe1
    project["aspice"] = aspice
    creator_store.save_project(project)
    return {"message": "已保存"}


@router.put("/projects/{project_id}/swe2")
async def save_swe2(project_id: str, req: SaveSwe2Request):
    """保存 SWE.2 数据。"""
    project = creator_store.get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    aspice = project.get("aspice") or creator_store._default_aspice()
    swe2 = aspice.get("swe2", {})
    update = req.model_dump(exclude_none=True)
    for k, v in update.items():
        swe2[k] = v
    aspice["swe2"] = swe2
    project["aspice"] = aspice
    creator_store.save_project(project)
    return {"message": "已保存"}


# ---- SWE.1 需求解析 ----

@router.post("/projects/{project_id}/swe1/parse-requirements")
async def parse_requirements(project_id: str, request: Request):
    """AI 解析客户原始需求（文本 + 可选多文件上传）→ 结构化需求项。

    文件支持 txt/docx/xlsx/pdf；AI 失败时规则兜底。
    手动读取 multipart，避免多文件字段在不同 FastAPI/Pydantic 版本下触发 422。
    """
    project = creator_store.get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    aspice = project.get("aspice") or creator_store._default_aspice()
    project_code = aspice.get("project_code", "") or project.get("name", "")

    form = await request.form()
    raw_text = str(form.get("text", "") or "")
    upload_files = []
    for field_name in ("files", "file"):
        for item in form.getlist(field_name):
            if hasattr(item, "filename") and hasattr(item, "read"):
                upload_files.append(item)

    # 文件解析
    for upload in upload_files:
        content = await upload.read()
        safe_name = os.path.basename(upload.filename or "")
        file_text = _extract_file_text(content, safe_name)
        if file_text:
            raw_text = (raw_text + "\n" + file_text).strip() if raw_text else file_text

    if not raw_text:
        raise HTTPException(status_code=400, detail="需求文本为空，请粘贴或上传客户需求")

    existing = aspice.get("swe1", {}).get("requirements", [])

    # AI 解析
    use_ai = True
    source = "rule"
    ai_error: Optional[str] = None
    req_items = []

    if use_ai:
        try:
            ai_service = _get_ai_service()
            if ai_service.enabled:
                ai_result = await ai_service.parse_requirements(raw_text, project_code)
                if ai_result and "_error" not in ai_result:
                    parsed = ai_result.get("requirements", [])
                    req_items = _assign_ids(parsed, project_code, existing)
                    source = "ai"
                else:
                    ai_error = (ai_result or {}).get("_error", "AI 解析失败")
        except Exception as e:
            ai_error = f"AI 调用异常: {str(e)}"

    # 规则兜底
    if not req_items:
        req_items = req_generator.generate_requirements(raw_text, project_code, existing)

    return {
        "source": source,
        "count": len(req_items),
        "requirements": req_items,
        "ai_error": ai_error,
    }


def _assign_ids(parsed: list, project_code: str, existing: list) -> list:
    """给 AI 解析出的需求项自动分配 OR ID + ReqID。"""
    if not parsed:
        return []
    existing_or_max = id_chain.next_or_seq(existing) - 1
    result = []
    for i, item in enumerate(parsed):
        if not isinstance(item, dict):
            continue
        or_seq = existing_or_max + i + 1
        or_id = id_chain.derive_or(project_code, or_seq)
        req_id = id_chain.derive_req(or_id, 1)
        priority = int(item.get("priority", 2)) if str(item.get("priority", "")).isdigit() else 2
        priority = min(max(priority, 1), 3)
        result.append({
            "or_id": or_id,
            "req_id": req_id,
            "software_mark": "原始",
            "content": str(item.get("content", "")),
            "or_desc": str(item.get("or_desc", "")),
            "sw_req_desc": str(item.get("sw_req_desc", "")),
            "category": _normalize_requirement_category(str(item.get("category", ""))),
            "asil": _normalize_asil(str(item.get("asil", "QM"))),
            "correctness": "Correct",
            "feasibility": "Feasible",
            "exception": "N/A",
            "milestone": str(item.get("milestone", "")),
            "owner": str(item.get("owner", "")),
            "input_source": "客户原始需求",
            "priority": priority,
            "actual_time": "NA",
            "release_version": "V1.0",
            "operation": str(item.get("operation", "")),
            "analysis": str(item.get("analysis", "")),
        })
    return result


def _normalize_requirement_category(category: str) -> str:
    """兼容旧 prompt 分类，归一化为 skill 枚举。"""
    mapping = {
        "Driver Basic Function": "Functional Requirements，Basic Functions",
        "Driver Safety Function": "Functional Requirements，Safety Requirements",
        "Driver Cybersecurity": "Functional Requirements，Cybersecurity Requirements",
        "Non-functional": "Non-Functional Requirements",
        "Non-camera": "Non-camera driver/tuning requirements",
        "基础功能": "Functional Requirements，Basic Functions",
        "安全需求": "Functional Requirements，Safety Requirements",
        "网络安全": "Functional Requirements，Cybersecurity Requirements",
        "非功能需求": "Non-Functional Requirements",
        "非Camera驱动": "Non-camera driver/tuning requirements",
    }
    options = {
        "Functional Requirements，Basic Functions",
        "Functional Requirements，Safety Requirements",
        "Functional Requirements，Cybersecurity Requirements",
        "Non-Functional Requirements",
        "Non-camera driver/tuning requirements",
    }
    if category in options:
        return category
    return mapping.get(category, "Functional Requirements，Basic Functions")


def _normalize_asil(asil: str) -> str:
    """归一化 ASIL 安全等级。"""
    value = asil.strip().upper().replace("ASIL-", "ASIL ")
    if value in {"QM", "N/A", "ASIL A", "ASIL B", "ASIL C", "ASIL D"}:
        return value
    if value in {"A", "B", "C", "D"}:
        return f"ASIL {value}"
    return "QM"


def _extract_file_text(content: bytes, filename: str) -> str:
    """从上传文件提取纯文本（txt/docx/xlsx）。"""
    import warnings
    warnings.filterwarnings("ignore")
    try:
        lower = filename.lower()
        if lower.endswith(".txt"):
            return content.decode("utf-8", errors="ignore")
        if lower.endswith(".docx"):
            import docx
            import io
            d = docx.Document(io.BytesIO(content))
            return "\n".join(p.text for p in d.paragraphs if p.text.strip())
        if lower.endswith((".xlsx", ".xls")):
            import openpyxl
            import io
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            lines = []
            for sn in wb.sheetnames:
                ws = wb[sn]
                for row in ws.iter_rows(values_only=True):
                    vals = [str(c) for c in row if c is not None]
                    if vals:
                        lines.append("\t".join(vals))
            return "\n".join(lines)
        if lower.endswith(".pdf"):
            try:
                import io
                import pdfplumber
                with pdfplumber.open(io.BytesIO(content)) as pdf:
                    return "\n".join(p.extract_text() or "" for p in pdf.pages)
            except ImportError:
                return ""
    except Exception:
        return ""
    return ""


# ---- SWE.2 映射生成 ----

@router.post("/projects/{project_id}/swe2/generate")
async def generate_swe2(project_id: str, use_ai: bool = Query(True)):
    """基于 SWE.1 需求自动生成 SWE.2 映射表（规则兜底 + AI 增强）。"""
    project = creator_store.get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    aspice = project.get("aspice") or creator_store._default_aspice()
    requirements = aspice.get("swe1", {}).get("requirements", [])
    if not requirements:
        raise HTTPException(status_code=400, detail="SWE.1 需求为空，请先完成需求分析")

    existing = aspice.get("swe2", {}).get("mappings", [])

    # 规则生成
    mappings = arch_generator.generate_mappings(requirements, existing=[])

    # AI 增强（可调整组件归属、补接口描述）
    source = "rule"
    ai_error: Optional[str] = None
    if use_ai:
        try:
            ai_service = _get_ai_service()
            if ai_service.enabled:
                ai_result = await ai_service.generate_arch_draft(requirements)
                if ai_result and "_error" not in ai_result:
                    ai_maps = ai_result.get("mappings", [])
                    if ai_maps:
                        mappings = _merge_ai_mappings(ai_maps, requirements, existing)
                        source = "ai"
                else:
                    ai_error = (ai_result or {}).get("_error", "AI 生成失败")
        except Exception as e:
            ai_error = f"AI 调用异常: {str(e)}"

    # 保存
    swe2 = aspice.get("swe2", {})
    swe2["mappings"] = mappings
    aspice["swe2"] = swe2
    project["aspice"] = aspice
    creator_store.save_project(project)

    return {
        "source": source,
        "count": len(mappings),
        "mappings": mappings,
        "ai_error": ai_error,
    }


def _component_id_from_swe2_id(swe2_id: str) -> str:
    """从 SWE.2 ID 末尾提取 A### 组件编号。"""
    import re
    m = re.search(r"-(A\d{3})$", swe2_id or "")
    return m.group(1) if m else ""


def _component_id_from_name(name: str) -> str:
    """按组件名反查 A###，兼容用户编辑后的映射表。"""
    text = (name or "").strip().lower()
    for cid, cname in arch_generator._COMPONENTS.items():
        if text == cname.lower() or cname.lower() in text:
            return cid
    return ""


def _default_swe2_components() -> list[dict]:
    """返回默认 8 组件属性，供 xlsx/docx 回灌共用。"""
    return [dict(c) for c in creator_store._default_aspice()["swe2"]["components"]]


def _build_swe2_mapping(
    swe1_id: str,
    swe2_id: str,
    component: str,
    description: str,
    release_version: str,
) -> Optional[dict]:
    """规范化一行 SWE.2 映射，兼容从 xlsx/docx 表格读取的空值。"""
    swe1_id = (swe1_id or "").strip()
    swe2_id = (swe2_id or "").strip()
    component = (component or "").strip()
    description = (description or "").strip()
    release_version = (release_version or "").strip()
    if not swe1_id and not swe2_id and not component:
        return None

    component_id = _component_id_from_swe2_id(swe2_id) or _component_id_from_name(component)
    if not component_id:
        component_id = "A004"
    if not component:
        component = arch_generator._COMPONENTS.get(component_id, "")
    if not swe2_id and swe1_id:
        swe2_id = f"{swe1_id}-{component_id}"

    return {
        "swe1_id": swe1_id,
        "swe2_id": swe2_id,
        "component_id": component_id,
        "component": component,
        "description": description,
        "release_version": release_version or "V1.0",
    }


def _dedupe_swe2_mappings(mappings: list[dict]) -> list[dict]:
    """按核心字段去重，兼容 docx 合并单元格造成的重复行。"""
    seen = set()
    unique = []
    for item in mappings:
        key = (
            item.get("swe1_id", ""),
            item.get("swe2_id", ""),
            item.get("component", ""),
            item.get("description", ""),
            item.get("release_version", ""),
        )
        if key in seen:
            continue
        seen.add(key)
        unique.append(item)
    return unique


def _parse_swe2_xlsx(path: str) -> tuple[list[dict], list[dict]]:
    """解析 SWE.2 导出的架构映射表 xlsx，返回 mappings/components。"""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    mappings: list[dict] = []
    if "需求映射" in wb.sheetnames:
        ws = wb["需求映射"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            mapping = _build_swe2_mapping(
                str(row[0] or "") if len(row) > 0 else "",
                str(row[1] or "") if len(row) > 1 else "",
                str(row[2] or "") if len(row) > 2 else "",
                str(row[3] or "") if len(row) > 3 else "",
                str(row[4] or "") if len(row) > 4 else "",
            )
            if mapping:
                mappings.append(mapping)

    default_components = _default_swe2_components()
    comp_by_name = {c["name"]: dict(c) for c in default_components}
    if "组件属性" in wb.sheetnames:
        ws = wb["组件属性"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            comp_name = str(row[0] or "").strip() if len(row) > 0 else ""
            attr = str(row[1] or "").strip() if len(row) > 1 else ""
            value = str(row[2] or "").strip() if len(row) > 2 else ""
            if not comp_name or comp_name not in comp_by_name:
                continue
            attr_key = attr.lower()
            if attr_key == "name" and value:
                comp_by_name[comp_name]["name"] = value
            elif attr_key in {"i2c address", "i2c地址", "i2c_addr"}:
                comp_by_name[comp_name]["i2c_addr"] = value
            elif attr_key in {"model", "型号"}:
                comp_by_name[comp_name]["model"] = value

    return _dedupe_swe2_mappings(mappings), list(comp_by_name.values())


def _cell_text(cell) -> str:
    """读取 docx 表格单元格文本，合并内部换行。"""
    return "\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip()).strip()


def _parse_swe2_docx(path: str) -> tuple[list[dict], list[dict]]:
    """解析 SWE.2 架构设计书 docx 中的需求映射表。"""
    from docx import Document

    doc = Document(path)
    mappings: list[dict] = []
    for table in doc.tables:
        if not table.rows:
            continue
        headers = [_cell_text(cell) for cell in table.rows[0].cells]
        normalized = [h.replace("\n", "").replace(" ", "") for h in headers]
        has_mapping_header = (
            any("SWE.1" in h for h in normalized)
            and any("SWE.2" in h for h in normalized)
            and any("ID含义" in h or "组件" in h for h in normalized)
        )
        if not has_mapping_header:
            continue

        def find_col(*names: str) -> int:
            for idx, header in enumerate(normalized):
                if any(name in header for name in names):
                    return idx
            return -1

        swe1_col = find_col("SWE.1")
        swe2_col = find_col("SWE.2")
        comp_col = find_col("ID含义", "组件")
        desc_col = find_col("备注", "描述")
        ver_col = find_col("发布版本", "版本")
        for row in table.rows[1:]:
            cells = [_cell_text(cell) for cell in row.cells]
            mapping = _build_swe2_mapping(
                cells[swe1_col] if 0 <= swe1_col < len(cells) else "",
                cells[swe2_col] if 0 <= swe2_col < len(cells) else "",
                cells[comp_col] if 0 <= comp_col < len(cells) else "",
                cells[desc_col] if 0 <= desc_col < len(cells) else "",
                cells[ver_col] if 0 <= ver_col < len(cells) else "",
            )
            if mapping:
                mappings.append(mapping)

    return _dedupe_swe2_mappings(mappings), _default_swe2_components()


@router.post("/projects/{project_id}/swe2/import")
async def import_swe2(project_id: str, file: UploadFile = File(...)):
    """回灌导入 SWE.2 架构映射表/架构设计书，写回 mappings/components 供继续编辑。"""
    project = creator_store.get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")

    safe_name = os.path.basename(file.filename or "swe2_import.xlsx")
    lower_name = safe_name.lower()
    if not lower_name.endswith((".xlsx", ".xls", ".docx")):
        raise HTTPException(status_code=400, detail="SWE.2 回灌支持 .xlsx/.xls 映射表或 .docx 架构设计书")

    out_dir = os.path.join("runtime", "creator_projects", project_id, "output")
    os.makedirs(out_dir, exist_ok=True)
    import_path = os.path.join(out_dir, f"导入SWE2_{safe_name}")
    content = await file.read()
    with open(import_path, "wb") as f:
        f.write(content)

    try:
        if lower_name.endswith(".docx"):
            mappings, components = _parse_swe2_docx(import_path)
        else:
            mappings, components = _parse_swe2_xlsx(import_path)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"SWE.2 文件解析失败: {str(e)}")

    if not mappings and not components:
        raise HTTPException(status_code=400, detail="未解析到 SWE.2 映射或组件属性，请检查文件格式")

    aspice = project.get("aspice") or creator_store._default_aspice()
    swe2 = aspice.get("swe2", {})
    if mappings:
        swe2["mappings"] = mappings
    if components:
        swe2["components"] = components
    swe2["current_step"] = max(int(swe2.get("current_step") or 0), 1)
    aspice["swe2"] = swe2
    project["aspice"] = aspice
    creator_store.save_project(project)

    return {
        "count": len(mappings),
        "component_count": len(components),
        "mappings": mappings,
        "components": components,
    }


def _merge_ai_mappings(ai_maps: list, requirements: list, existing: list) -> list:
    """把 AI 返回的映射建议转换为完整 mappings（带 swe2_id）。"""
    from app.core.arch_generator import _COMPONENTS
    merged = list(existing)
    for am in ai_maps:
        if not isinstance(am, dict):
            continue
        req_id = am.get("req_id", "")
        comp_ids = am.get("component_ids", [])
        iface = am.get("interface_desc", "")
        for comp_id in comp_ids:
            comp_id = comp_id if comp_id.startswith("A") else f"A{comp_id}"
            a_seq = id_chain.next_arch_seq(merged, req_id)
            swe2_id = id_chain.derive_arch(req_id, a_seq)
            merged.append({
                "swe1_id": req_id,
                "swe2_id": swe2_id,
                "component_id": comp_id,
                "component": _COMPONENTS.get(comp_id, ""),
                "description": iface,
                "release_version": "V1.0",
            })
    return merged


# ---- 导出 ----

@router.post("/projects/{project_id}/swe1/export")
async def export_swe1(project_id: str):
    """导出 SWE.1：需求说明书 docx + 需求详细表 xlsx。

    经 aspice_adapter 转 skill project.json，调 build_docx / build_excel
    填充双语四小节 SRS docx 模板 + 24 列 7 sheet xlsx 模板。
    """
    project = creator_store.get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    out_dir = os.path.join("runtime", "creator_projects", project_id, "output")
    os.makedirs(out_dir, exist_ok=True)

    code = (project.get("aspice") or {}).get("project_code", "") or project.get("name", "")
    safe_code = code.replace("/", "_").replace("\\", "_")
    docx_path = os.path.join(out_dir, f"{safe_code}_软件需求说明书.docx")
    xlsx_path = os.path.join(out_dir, f"{safe_code}_需求详细表.xlsx")

    try:
        req_generator.export_swe1_docx(project, docx_path)
        req_generator.export_swe1_xlsx(project, xlsx_path)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"导出失败: {str(e)}")

    return {
        "docx": os.path.basename(docx_path),
        "xlsx": os.path.basename(xlsx_path),
    }


# ---- 校验 ----

@router.post("/projects/{project_id}/swe1/validate")
async def validate_swe1(project_id: str):
    """校验 SWE.1：经 aspice_adapter 转 skill project.json 后调 aspice_validate。

    V01-V18 校验：ID 格式/连续/唯一、Input↔inputs 一致、Deleted 话术、
    SRS 覆盖率、双语完整等。返回 errors/warnings/passed/report。
    """
    project = creator_store.get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")

    from app.core import aspice_adapter
    from app.core import aspice_validate

    skill_json = aspice_adapter.to_skill_project(project)

    # 可选：对已导出的 xlsx/docx 做往返校验（V14/V15）
    out_dir = os.path.join("runtime", "creator_projects", project_id, "output")
    code = (project.get("aspice") or {}).get("project_code", "") or project.get("name", "")
    safe_code = code.replace("/", "_").replace("\\", "_")
    excel_path = os.path.join(out_dir, f"{safe_code}_需求详细表.xlsx")
    docx_path = os.path.join(out_dir, f"{safe_code}_软件需求说明书.docx")
    excel_path = excel_path if os.path.isfile(excel_path) else None
    docx_path = docx_path if os.path.isfile(docx_path) else None

    try:
        result = aspice_validate.validate_project(
            skill_json, excel_path=excel_path, docx_path=docx_path,
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"校验失败: {str(e)}")

    return result


@router.post("/projects/{project_id}/swe2/export")
async def export_swe2(project_id: str):
    """导出 SWE.2：架构&概要设计书 docx + 映射表 xlsx。"""
    project = creator_store.get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    out_dir = os.path.join("runtime", "creator_projects", project_id, "output")
    os.makedirs(out_dir, exist_ok=True)

    code = (project.get("aspice") or {}).get("project_code", "") or project.get("name", "")
    safe_code = code.replace("/", "_").replace("\\", "_")
    docx_path = os.path.join(out_dir, f"{safe_code}_软件架构&概要设计书.docx")
    xlsx_path = os.path.join(out_dir, f"{safe_code}_架构映射表.xlsx")

    try:
        arch_generator.export_swe2_docx(project, docx_path)
        arch_generator.export_swe2_xlsx(project, xlsx_path)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"导出失败: {str(e)}")

    return {
        "docx": os.path.basename(docx_path),
        "xlsx": os.path.basename(xlsx_path),
    }


# ---- 下载 ----

@router.get("/projects/{project_id}/download/{filename}")
async def download_file(project_id: str, filename: str):
    """下载导出的文件。"""
    path = creator_store.get_download_path(project_id, filename)
    if not path:
        raise HTTPException(status_code=404, detail="文件不存在")
    return FileResponse(path, filename=filename, media_type="application/octet-stream")
