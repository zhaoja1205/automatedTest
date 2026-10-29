"""数据统计页 AI 文案：数据汇总、Prompt、Excel 写入与服务解析。"""
import asyncio
import os
import sys
import tempfile
from types import SimpleNamespace

_HERE = os.path.dirname(os.path.abspath(__file__))
_BACKEND_DIR = os.path.dirname(_HERE)
sys.path.insert(0, _BACKEND_DIR)

from openpyxl import Workbook, load_workbook  # noqa: E402

from app.ai.cache import AICache  # noqa: E402
from app.ai.providers.base import LLMResponse  # noqa: E402
from app.ai.service import AIService  # noqa: E402
from app.core.excel_handler import ExcelHandler  # noqa: E402
from app.core.test_case import TestCase, WorkspaceConfig, SSHConfig  # noqa: E402
from app.core.execution_state import ExecutionTask  # noqa: E402


def _build_statistics_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "功能测试"
    ws["B13"] = "测试用例ID"
    ws["C13"] = "测试类型"
    ws["F13"] = "测试用例描述"

    stats = wb.create_sheet("数据统计")
    stats["B2"] = "测试结果汇总(Round1)"
    stats["B3"] = "测试分类"
    stats["B12"] = "测试基本信息"
    stats["B13"] = "质量/效率目标"
    stats["B14"] = "测试总结"
    stats.merge_cells("C12:G12")
    stats.merge_cells("C13:G13")
    stats.merge_cells("C14:G14")
    wb.save(path)
    return path


def test_fill_statistics_summary_writes_merged_cells(tmp_path):
    path = _build_statistics_workbook(tmp_path / "case.xlsx")
    handler = ExcelHandler()
    handler.load(str(path))

    assert handler.find_statistics_sheet_name() == "数据统计"
    filled = handler.fill_statistics_summary(
        basic_info="测试人员：张三\n测试版本：V1.0",
        quality_goal="执行率目标 100%，实际 95.0%",
        summary="本轮测试整体通过，遗留 1 个一般问题。",
    )
    assert filled == ["测试基本信息", "质量/效率目标", "测试总结"]

    handler.save(str(path))
    ws = load_workbook(str(path), data_only=False)["数据统计"]
    assert ws["C12"].value == "测试人员：张三\n测试版本：V1.0"
    assert ws["C13"].value == "执行率目标 100%，实际 95.0%"
    assert ws["C14"].value == "本轮测试整体通过，遗留 1 个一般问题。"
    assert ws["C12"].alignment.wrap_text is True
    assert ws["C12"].alignment.vertical == "top"


def test_fill_statistics_summary_missing_sheet(tmp_path):
    path = _build_statistics_workbook(tmp_path / "case.xlsx")
    wb = load_workbook(str(path))
    del wb["数据统计"]
    wb.save(str(path))

    handler = ExcelHandler()
    handler.load(str(path))
    assert handler.find_statistics_sheet_name() == ""
    assert handler.fill_statistics_summary(basic_info="x") == []


class _FakeProvider:
    def __init__(self, content):
        self.content = content
        self.calls = []

    async def complete(self, system_prompt, user_prompt, **kwargs):
        self.calls.append({"system": system_prompt, "user": user_prompt, **kwargs})
        return LLMResponse(content=self.content, model="fake-model", input_tokens=20, output_tokens=30)


def test_generate_test_statistics_parses_json():
    payload = '{"basic_info": "测试人员：张三\\n测试版本：V1.0", "quality_goal": "1. 执行率 100%", "summary": "整体通过。"}'
    svc = AIService({"ai_enabled": True, "ai_protocol": "openai", "ai_api_key": "k"})
    svc._provider = _FakeProvider(payload)
    svc._cache = AICache(cache_dir=tempfile.mkdtemp(prefix="ai_stats_cache_"))

    result = asyncio.run(svc.generate_test_statistics({
        "workbook": "case", "tester": "张三", "version": "V1.0",
        "total": 10, "executed": 8, "pass": 7, "fail": 1,
    }))
    assert result["basic_info"] == "测试人员：张三\n测试版本：V1.0"
    assert result["quality_goal"] == "1. 执行率 100%"
    assert result["summary"] == "整体通过。"
    assert result["_model"] == "fake-model"
    assert len(svc._provider.calls) == 1


def test_generate_test_statistics_disabled():
    svc = AIService({"ai_enabled": False})
    result = asyncio.run(svc.generate_test_statistics({"total": 1}))
    assert "_error" in result


def test_build_statistics_data_counts_statuses_and_rates():
    from app.api.routes import _build_statistics_data

    cases = [
        {"case_id": "FT_001", "source_sheet": "功能测试", "status": "Pass"},
        {"case_id": "FT_002", "source_sheet": "功能测试", "status": "Fail",
         "description": "起流测试", "ai_reason": "输出偏绿", "bug_severity": "严重"},
        {"case_id": "FT_003", "source_sheet": "功能测试", "status": "Block"},
        {"case_id": "FT_004", "source_sheet": "故障测试", "status": "NA"},
        {"case_id": "FT_005", "source_sheet": "故障测试", "status": "NT"},
        {"case_id": "FT_006", "source_sheet": "故障测试", "status": "Review"},
    ]
    session = SimpleNamespace(
        test_cases=cases,
        workspace=WorkspaceConfig(tester_name="张三", test_version="V1.0"),
        current_task=ExecutionTask(
            started_at="2026-10-29T10:00:00", ended_at="2026-10-29T10:12:00",
        ),
        ssh_config=SSHConfig(host="192.168.1.10", username="root"),
        excel_path="uploads/demo/case.xlsx",
    )
    handler = SimpleNamespace(file_path="", test_cases=[])

    data = _build_statistics_data(session, handler)
    assert data["total"] == 6
    assert data["pass"] == 1
    assert data["fail"] == 1
    assert data["block"] == 1
    assert data["na"] == 1
    assert data["nt"] == 1
    assert data["review"] == 1
    # 执行数 = 总数 - NT - NA
    assert data["executed"] == 4
    assert data["execute_rate"] == round(4 / 6 * 100, 1)
    assert data["pass_rate"] == 25.0
    assert data["severity_dist"] == {"严重": 1}
    assert data["workbook"] == "case"
    assert data["environment"] == "root@192.168.1.10:22"
    assert data["duration_seconds"] == 720.0
    assert data["test_date"] == "2026-10-29"
    assert data["failures"][0]["case_id"] == "FT_002"
    assert data["failures"][0]["reason"] == "输出偏绿"


def test_build_statistics_data_falls_back_to_loaded_cases():
    from app.api.routes import _build_statistics_data

    handler = SimpleNamespace(
        file_path="exports/case.xlsx",
        test_cases=[
            TestCase(case_id="FT_001", source_sheet="功能测试", status="Pass"),
            TestCase(case_id="FT_002", source_sheet="功能测试", status="Fail",
                     ai_reason="输出异常"),
        ],
    )
    session = SimpleNamespace(
        test_cases=[],
        workspace=WorkspaceConfig(),
        current_task=ExecutionTask(),
        ssh_config=SSHConfig(),
        excel_path="",
    )
    data = _build_statistics_data(session, handler)
    assert data["total"] == 2
    assert data["pass"] == 1
    assert data["fail"] == 1
    assert data["workbook"] == "case"
    assert data["test_date"]


def test_download_results_fills_statistics_with_ai(tmp_path):
    from types import SimpleNamespace
    from app.api import routes as routes_module

    path = _build_statistics_workbook(tmp_path / "case.xlsx")
    cases = [
        {"case_id": "FT_001", "source_sheet": "功能测试", "status": "Pass"},
        {"case_id": "FT_002", "source_sheet": "功能测试", "status": "Fail",
         "description": "起流测试", "ai_reason": "输出偏绿", "bug_severity": "严重"},
    ]
    session = SimpleNamespace(
        excel_path=str(path),
        test_cases=cases,
        workspace=WorkspaceConfig(tester_name="张三", test_version="V1.0", image_download_path=""),
        current_task=ExecutionTask(
            started_at="2026-10-29T10:00:00", ended_at="2026-10-29T10:05:00",
        ),
        ssh_config=SSHConfig(host="192.168.1.10", username="root"),
    )
    request = SimpleNamespace(state=SimpleNamespace(session=session))

    class _Svc:
        enabled = True

        async def generate_test_statistics(self, data):
            assert data["executed"] == 2
            return {
                "basic_info": "测试人员：张三\n测试版本：V1.0",
                "quality_goal": "执行率目标 100%，实际 100%",
                "summary": "本轮共 2 条用例，Fail 1 条，需修复后回归。",
            }

    original = routes_module._get_ai_service
    routes_module._get_ai_service = lambda _session: _Svc()
    try:
        response = asyncio.run(routes_module.download_results(request, include_log=True))
        out_path = response.path
        ws = load_workbook(out_path, data_only=False)["数据统计"]
        assert ws["C12"].value == "测试人员：张三\n测试版本：V1.0"
        assert ws["C13"].value == "执行率目标 100%，实际 100%"
        assert ws["C14"].value == "本轮共 2 条用例，Fail 1 条，需修复后回归。"
    finally:
        routes_module._get_ai_service = original


def test_download_results_survives_ai_failure(tmp_path):
    from types import SimpleNamespace
    from app.api import routes as routes_module

    path = _build_statistics_workbook(tmp_path / "case.xlsx")
    session = SimpleNamespace(
        excel_path=str(path),
        test_cases=[
            {"case_id": "FT_001", "source_sheet": "功能测试", "status": "Pass"},
        ],
        workspace=WorkspaceConfig(image_download_path=""),
        current_task=ExecutionTask(),
        ssh_config=SSHConfig(),
    )
    request = SimpleNamespace(state=SimpleNamespace(session=session))

    class _BrokenSvc:
        enabled = True

        async def generate_test_statistics(self, data):
            raise RuntimeError("模拟 AI 超时")

    original = routes_module._get_ai_service
    routes_module._get_ai_service = lambda _session: _BrokenSvc()
    try:
        response = asyncio.run(routes_module.download_results(request, include_log=True))
        assert os.path.exists(response.path)
        ws = load_workbook(response.path, data_only=False)["数据统计"]
        assert ws["C12"].value is None
    finally:
        routes_module._get_ai_service = original
