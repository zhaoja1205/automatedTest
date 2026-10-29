"""BugList统计 页解析单测：页签识别与行数据同步。"""
import datetime
import os
import sys

_HERE = os.path.dirname(os.path.abspath(__file__))
_BACKEND_DIR = os.path.dirname(_HERE)
sys.path.insert(0, _BACKEND_DIR)

from openpyxl import Workbook  # noqa: E402

from app.core.excel_handler import ExcelHandler  # noqa: E402


def _build_workbook(path, bug_rows=None, with_bug_sheet=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "功能测试"
    ws["B13"] = "测试用例ID"
    ws["C13"] = "测试类型"
    ws["F13"] = "测试用例描述"
    ws["B15"] = "FT_001"
    ws["C15"] = "功能"
    ws["F15"] = "起流测试"
    ws["H15"] = "执行命令"
    ws["I15"] = "输出正常"

    if with_bug_sheet:
        bug_ws = wb.create_sheet("BugList统计")
        headers = [
            "No.", "Bug ID", "Module", "Description", "Version",
            "Severity", "Status", "Tester", "Date", "Effective", "Comment",
        ]
        for col, title in enumerate(headers, start=2):
            bug_ws.cell(row=2, column=col).value = title
        for offset, values in enumerate(bug_rows or []):
            for col, value in enumerate(values, start=2):
                bug_ws.cell(row=3 + offset, column=col).value = value

    wb.save(path)
    return path


def test_is_buglist_sheet_matches_template_names():
    assert ExcelHandler.is_buglist_sheet("BugList统计")
    assert ExcelHandler.is_buglist_sheet("Bug List")
    assert ExcelHandler.is_buglist_sheet("缺陷列表")
    assert not ExcelHandler.is_buglist_sheet("功能测试")
    assert not ExcelHandler.is_buglist_sheet("缺陷测试")


def test_load_bug_list_parses_rows_and_formats_date(tmp_path):
    path = _build_workbook(
        tmp_path / "case.xlsx",
        bug_rows=[[
            1, "BUG-001", "ISP", "图像偏色", "V1.0",
            "Major", "Open", "张三", datetime.datetime(2026, 10, 29),
            "是", "待确认",
        ]],
    )
    handler = ExcelHandler()
    handler.load(str(path))

    assert handler.bug_sheet_name == "BugList统计"
    assert handler.bug_list == [{
        "no": "1",
        "bug_id": "BUG-001",
        "module": "ISP",
        "description": "图像偏色",
        "version": "V1.0",
        "severity": "Major",
        "status": "Open",
        "tester": "张三",
        "date": "2026-10-29",
        "effective": "是",
        "comment": "待确认",
        "row_number": 3,
        "source_sheet": "BugList统计",
    }]


def test_load_bug_list_empty_without_bug_sheet(tmp_path):
    path = _build_workbook(tmp_path / "case.xlsx", with_bug_sheet=False)
    handler = ExcelHandler()
    handler.load(str(path))

    assert handler.bug_sheet_name == ""
    assert handler.bug_list == []


def test_sync_failed_cases_only_registers_fail_and_dedupes():
    from app.core.bug_list import sync_failed_cases_to_bug_list
    from app.core.test_case import TestCase

    failed = TestCase(
        case_id="FT_001", source_sheet="功能测试", test_type="功能",
        description="起流测试", tester="张三", test_version="V1.0",
        test_date="2026-10-29 15:00:00", status="Fail",
        ai_reason="帧率只有 12fps，低于预期",
    )
    passed = TestCase(case_id="FT_002", source_sheet="功能测试", status="Pass")

    entries, added = sync_failed_cases_to_bug_list([failed, passed], [])
    assert added == 1
    assert len(entries) == 1
    assert entries[0]["bug_id"] == "BUG-FT_001"
    assert entries[0]["status"] == "Open"
    # 未启用 AI 时 Severity 使用默认等级
    assert entries[0]["severity"] == "待评估"
    assert entries[0]["effective"] == "待确认"
    assert entries[0]["date"] == "2026-10-29"
    assert "12fps" in entries[0]["description"]
    assert entries[0]["comment"] == ""

    # 同一用例再次失败时更新原条目，不新增重复 Bug
    entries2, added2 = sync_failed_cases_to_bug_list([failed], entries)
    assert added2 == 0
    assert len(entries2) == 1


def test_save_bug_list_writes_entries_back_to_sheet(tmp_path):
    from app.core.bug_list import sync_failed_cases_to_bug_list
    from app.core.test_case import TestCase

    path = _build_workbook(tmp_path / "case.xlsx")
    handler = ExcelHandler()
    handler.load(str(path))

    failed = TestCase(
        case_id="FT_001", source_sheet="功能测试", test_type="功能",
        description="起流测试", status="Fail", ai_reason="输出异常",
    )
    entries, _ = sync_failed_cases_to_bug_list([failed], handler.bug_list)
    handler.bug_list = entries
    handler.save_bug_list()
    handler.save(str(path))

    reloaded = ExcelHandler()
    reloaded.load(str(path))
    assert reloaded.bug_sheet_name == "BugList统计"
    assert reloaded.bug_list[0]["bug_id"] == "BUG-FT_001"
    assert reloaded.bug_list[0]["row_number"] == 3
